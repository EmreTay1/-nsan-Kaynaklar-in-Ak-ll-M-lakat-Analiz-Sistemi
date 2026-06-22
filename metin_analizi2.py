# -*- coding: utf-8 -*-
import requests
from docx import Document
import os
import json

# --- KONFİGÜRASYON ---
# Yerel LM Studio sunucunuzun çalıştığından emin olun.
LM_STUDIO_API_URL = "http://localhost:1234/v1/chat/completions"
# İndirdiğiniz GGUF model dosyasının adını doğrulayın.
MODEL_NAME = "mistral-7b-instruct-v0.3.Q4_0"
# Mülakat transkriptini içeren girdi DOCX dosyası.
INPUT_DOCX_FILE = "duygu_analizli_transkript.docx"
# Analizin kaydedileceği çıktı DOCX dosyası.
OUTPUT_DOCX_FILE = "analiz_sonucu_yeni_sablon.docx"


def read_text_from_docx(file_path):
    """Bir .docx dosyasındaki tüm metni okur."""
    try:
        doc = Document(file_path)
        full_text = [para.text for para in doc.paragraphs]
        return '\n'.join(full_text)
    except Exception as e:
        print(f"HATA: '{file_path}' dosyası okunurken bir hata oluştu: {e}")
        return None


def get_llm_analysis(prompt, model_name):
    """LM Studio API aracılığıyla yerel LLM'e bir prompt gönderir ve yanıtı alır."""
    payload = {
        "model": model_name,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.7,
        "max_tokens": 2000
    }
    try:
        response = requests.post(LM_STUDIO_API_URL, json=payload)
        response.raise_for_status()
        response_data = response.json()

        print("API Yanıtı:", json.dumps(response_data, indent=2, ensure_ascii=False))

        if "choices" in response_data and response_data["choices"]:
            return response_data["choices"][0]["message"]["content"]
        else:
            print("HATA: API yanıtında 'choices' anahtarı bulunamadı veya boş.")
            return None
    except requests.exceptions.ConnectionError:
        print("HATA: LM Studio API sunucusuna bağlanılamadı. Lütfen LM Studio'yu kontrol edin.")
        return None
    except requests.exceptions.HTTPError as http_err:
        print(f"HATA: HTTP hatası: {http_err}")
        print(f"Yanıt İçeriği: {response.text}")
        return None
    except Exception as e:
        print(f"HATA: Analiz sırasında beklenmedik bir hata oluştu: {e}")
        return None


def analyze_interview():
    """Mülakat analiz süreci."""
    print(f"'{INPUT_DOCX_FILE}' dosyasındaki mülakat metni okunuyor...")
    interview_text = read_text_from_docx(INPUT_DOCX_FILE)
    if not interview_text or not interview_text.strip():
        print(f"HATA: '{INPUT_DOCX_FILE}' dosyası boş veya okunamadı.")
        return

    print("Mülakat metni başarıyla okundu.")

    # --- YENİ İK ANALİZ PROMPT'U ---
    print(f"'{MODEL_NAME}' modeline İK analizi isteği gönderiliyor...")

    prompt_analysis = f"""
    Aşağıdaki mülakat metnini profesyonel bir İnsan Kaynakları (İK) uzmanı gibi analiz et.
    Analizini, metnin dili ne olursa olsun tamamen Türkçe olarak oluştur.
    Analizi yalnızca aşağıdaki başlık yapısına göre oluştur. Başlıklar dışına çıkma, ekstra yorum ekleme.

    -----------------------------
    🧩 MÜLAKAT ANALİZ RAPORU
    -----------------------------

    ### 1. Aday Değerlendirme Puanlama Tablosu
    Adayın performansını aşağıdaki kriterlere göre 1–5 arasında puanla.
    Her kriter için şu formatı kullan:
    • Kriter Adı: 4/5 – Kısa gerekçe.

    Kriterler:
    - İletişim Becerisi
    - Motivasyon ve Tutku
    - Analitik / Düşünsel Beceriler
    - Profesyonel Tutum
    - Geçmiş Deneyim Uyumu
    - Liderlik ve Girişimcilik
    - Zayıflıklarla Başa Çıkma Yetisi
    - Uzun Vadeli Potansiyel
    - Genel Etki / İzlenim

    Genel Ortalama Puan: (1–5 arasında aritmetik ortalama değeri)

    ### 2. Beceriler (Teknik ve Mesleki Yetkinlikler)
    Mülakat metninden çıkarılabilen teknik, yazılım, mesleki veya operasyonel becerileri listele.
    • Python – Veri analizi ve otomasyon deneyimi.
    • Proje Yönetimi – Takım koordinasyonu, zaman planlama.

    ### 3. Davranışsal Gözlemler
    Adayın iletişim tarzı, özgüveni, stresle başa çıkma biçimi ve takım çalışmasına yatkınlığına dair çıkarımları yaz.
    • Örnek: Kendine güveni yüksek, ancak bazen savunmacı tavır sergiliyor.

    ### 4. Rol Uyumu Tahmini
    Adayın kişilik, beceri ve tutumuna göre hangi tür rol/pozisyon yapısına daha uygun olduğunu belirt.
    • Örnek: Analitik karar destek veya proje koordinasyonu pozisyonları için uygun.

    --- MÜLAKAT METNİ ---
    {interview_text}

    --- ANALİZİ BURADAN BAŞLAT ---
    """

    analysis_result = get_llm_analysis(prompt_analysis, MODEL_NAME)
    if not analysis_result:
        print("Analiz alınamadı. İşlem durduruluyor.")
        return

    print("Analiz başarıyla tamamlandı, DOCX dosyasına kaydediliyor...")

    # --- SONUÇLARI DOCX'E YAZ ---
    try:
        doc = Document()
        doc.add_heading('Mülakat Analizi Sonucu', level=0)
        for line in analysis_result.split('\n'):
            if line.strip():
                doc.add_paragraph(line.strip())

        doc.save(OUTPUT_DOCX_FILE)
        print(f"Analiz sonuçları '{OUTPUT_DOCX_FILE}' dosyasına başarıyla kaydedildi.")
        # os.startfile(OUTPUT_DOCX_FILE)  # Windows'ta dosyayı otomatik açmak için
    except Exception as e:
        print(f"HATA: Çıktı dosyası kaydedilirken bir hata oluştu: {e}")
        with open("hata_ayiklama_cikti.txt", "w", encoding="utf-8") as f:
            f.write(analysis_result)
        print("Ham analiz çıktısı 'hata_ayiklama_cikti.txt' dosyasına kaydedildi.")


if __name__ == "__main__":
    analyze_interview()
